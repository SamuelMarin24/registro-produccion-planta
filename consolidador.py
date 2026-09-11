"""
CONSOLIDADOR  ·  Registro de Producción en Planta
=================================================

Corre SOLO en el PC dedicado (siempre encendido).

Qué hace, en bucle, cada CICLO_SEGUNDOS:
  1. Mira la carpeta PENDIENTES (donde las máquinas dejan sus archivos .json).
  2. Toma los archivos, los ordena por fecha de llegada y los agrega como filas
     al Excel consolidado.
  3. Mueve cada archivo ya procesado a la carpeta PROCESADOS (respaldo).

Por qué esto es seguro:
  - Este programa es el ÚNICO que abre y escribe el Excel. Las máquinas nunca
    lo tocan, así que no puede corromperse por escrituras simultáneas.
  - Si el Excel está abierto por alguien, la escritura falla y los archivos se
    quedan en PENDIENTES; se procesan en el siguiente ciclo. No se pierde nada.

Cómo se ejecuta en el PC dedicado:
  - Doble clic en  consolidador.bat   (queda una ventana negra abierta con el log).
  - O configurarlo como tarea programada de Windows al iniciar sesión.

Detener: cerrar la ventana, o Ctrl + C.
"""

import os
import json
import time
import shutil
import logging
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment

# Reutilizamos la definición de columnas, el encabezado y la config del proyecto,
# para que el Excel salga EXACTAMENTE igual que antes.
from data_manager import COLUMNAS_CONSOLIDADO, DataManager
from config import Config, CICLO_SEGUNDOS

# ── Parámetros ────────────────────────────────────────────────────────────────
REINTENTOS_EXCEL = 3         # reintentos si el Excel está ocupado momentáneamente
ESPERA_REINTENTO = 5         # segundos entre reintentos

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [Consolidador] %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler(
            os.path.join(os.path.dirname(__file__), "consolidador.log"),
            encoding="utf-8"
        ),
    ],
)


class Consolidador:

    def __init__(self):
        self.config = Config()
        # DataManager nos da _write_header y get_maquina_info ya hechos.
        # No pasa nada si las maestras no cargan: solo usamos utilidades.
        try:
            self.dm = DataManager.__new__(DataManager)
            self.dm.config = self.config
            self.dm.maestras = {}
            self.dm.ops = {}
            import threading
            self.dm._lock = threading.Lock()
        except Exception as e:
            logging.warning(f"No se pudo inicializar utilidades de DataManager: {e}")
            self.dm = None

        self.ruta_excel  = self.config.get_ruta_excel()
        self.pendientes  = self.config.get_carpeta_pendientes()
        self.procesados  = self.config.get_carpeta_procesados()

    # ── Utilidades de Excel ─────────────────────────────────────────────────

    def _abrir_o_crear_libro(self):
        """Abre el Excel consolidado, o lo crea con encabezado si no existe."""
        if os.path.exists(self.ruta_excel):
            wb = openpyxl.load_workbook(self.ruta_excel)
            ws = wb.active
            if ws.max_row == 0 or ws.cell(1, 1).value is None:
                ws.title = "Consolidado"
                self._escribir_header(ws)
            return wb, ws
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Consolidado"
        self._escribir_header(ws)
        return wb, ws

    def _escribir_header(self, ws):
        if self.dm is not None:
            self.dm._write_header(ws)
        else:
            for i, col in enumerate(COLUMNAS_CONSOLIDADO, 1):
                ws.cell(row=1, column=i, value=col)

    # ── Procesamiento de un ciclo ───────────────────────────────────────────

    def _archivos_pendientes(self):
        if not os.path.isdir(self.pendientes):
            return []
        try:
            archivos = [a for a in os.listdir(self.pendientes)
                        if a.lower().endswith(".json")]
        except Exception as e:
            logging.warning(f"No se pudo listar PENDIENTES: {e}")
            return []
        # Ordenar por fecha de modificación → respeta el orden de llegada
        rutas = [os.path.join(self.pendientes, a) for a in archivos]
        rutas.sort(key=lambda p: os.path.getmtime(p))
        return rutas

    def _procesar_ciclo(self):
        rutas = self._archivos_pendientes()
        if not rutas:
            return 0

        # Leer todos los registros válidos
        registros = []
        leidos = []
        for ruta in rutas:
            try:
                with open(ruta, "r", encoding="utf-8") as f:
                    registros.append(json.load(f))
                leidos.append(ruta)
            except Exception as e:
                logging.warning(f"Archivo ilegible, lo dejo para revisar: "
                                f"{os.path.basename(ruta)} ({e})")

        if not registros:
            return 0

        # Escribir al Excel con reintentos (por si está ocupado un instante)
        for intento in range(1, REINTENTOS_EXCEL + 1):
            try:
                wb, ws = self._abrir_o_crear_libro()
                for reg in registros:
                    fila = [reg.get(col, "") for col in COLUMNAS_CONSOLIDADO]
                    ws.append(fila)
                    for cell in ws[ws.max_row]:
                        cell.font = Font(name="Calibri", size=10)
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                os.makedirs(os.path.dirname(self.ruta_excel), exist_ok=True)
                wb.save(wb_path := self.ruta_excel)
                wb.close()
                break
            except PermissionError:
                logging.warning(f"Excel ocupado (¿abierto por alguien?). "
                                f"Intento {intento}/{REINTENTOS_EXCEL}.")
                if intento == REINTENTOS_EXCEL:
                    logging.warning("No se pudo escribir el Excel este ciclo. "
                                    "Los archivitos quedan en PENDIENTES.")
                    return 0
                time.sleep(ESPERA_REINTENTO)
            except Exception as e:
                logging.error(f"Error escribiendo el Excel: {e}. "
                              "Los archivitos quedan en PENDIENTES.")
                return 0

        # Mover a PROCESADOS, ordenado por subcarpeta mensual: "Mayo 2026", etc.
        # Así la carpeta no se llena de cientos de miles de archivos sueltos
        # y borrar lo viejo es tan fácil como borrar la subcarpeta del mes.
        meses_es = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
        hoy = datetime.now()
        subcarpeta = f"{meses_es[hoy.month - 1]} {hoy.year}"
        destino_mes = os.path.join(self.procesados, subcarpeta)
        os.makedirs(destino_mes, exist_ok=True)

        movidos = 0
        for ruta in leidos:
            try:
                destino = os.path.join(destino_mes, os.path.basename(ruta))
                # Si por alguna razón ya existe, le añadimos sufijo
                if os.path.exists(destino):
                    base, ext = os.path.splitext(destino)
                    destino = f"{base}_{datetime.now().strftime('%H%M%S')}{ext}"
                shutil.move(ruta, destino)
                movidos += 1
            except Exception as e:
                logging.warning(f"No se pudo mover {os.path.basename(ruta)} "
                                f"a PROCESADOS: {e}")
        return movidos

    # ── Bucle principal ──────────────────────────────────────────────────────

    def run(self):
        logging.info("=" * 60)
        logging.info("Consolidador iniciado.")
        logging.info(f"  PENDIENTES : {self.pendientes}")
        logging.info(f"  PROCESADOS : {self.procesados}")
        logging.info(f"  EXCEL      : {self.ruta_excel}")
        logging.info(f"  Ciclo      : cada {CICLO_SEGUNDOS} s")
        logging.info("=" * 60)
        while True:
            try:
                n = self._procesar_ciclo()
                if n:
                    logging.info(f"{n} registro(s) agregado(s) al Excel.")
            except Exception as e:
                logging.error(f"Error inesperado en el ciclo: {e}")
            time.sleep(CICLO_SEGUNDOS)


if __name__ == "__main__":
    try:
        Consolidador().run()
    except KeyboardInterrupt:
        logging.info("Consolidador detenido por el usuario.")
